"""Tailored CV + cover letter generation with an independent fact-check pass."""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document

from .. import db
from ..ai import AIError, AIUnavailable, BudgetExceeded, ClaudeService
from ..config import OUTPUT_ROOT, Config
from ..events import log_event
from ..models import MaterialsVerification, TailoredMaterials
from ..pipeline.normalize import ascii_lower
from ..profile import CandidateProfile


def _slug(value: str | None, limit: int = 40) -> str:
    return re.sub(r"[^a-z0-9]+", "-", ascii_lower(value)).strip("-")[:limit] or "unknown"


def _feedback(verification: MaterialsVerification) -> str:
    return "\n".join(f"- [{c.location}] \"{c.claim}\": {c.reason}" for c in verification.unsupported_claims)


def generate_for_opportunity(opportunity_id: int, run_id: int | None, ai: ClaudeService, cfg: Config, profile: CandidateProfile) -> dict[str, Any]:
    row = db.fetch_one("SELECT * FROM opportunities WHERE id = %s", (opportunity_id,))
    if row is None:
        raise LookupError(f"Opportunity {opportunity_id} not found")
    job = dict(row)
    analysis = row.get("analysis") or {}
    summary = {"id": opportunity_id, "title": row["title"]}

    try:
        materials = ai.generate_materials(job, analysis, run_id)
        verification = ai.verify_materials(job, materials, run_id)
        if not verification.all_claims_supported or verification.unsupported_claims:
            materials = ai.generate_materials(job, analysis, run_id, feedback=_feedback(verification))
            verification = ai.verify_materials(job, materials, run_id)
    except (BudgetExceeded, AIUnavailable) as exc:
        db.execute("UPDATE opportunities SET materials_status = 'none', updated_at = now() WHERE id = %s", (opportunity_id,))
        return {**summary, "outcome": "materials_deferred", "error": str(exc)}
    except AIError as exc:
        db.execute("UPDATE opportunities SET materials_status = 'failed', updated_at = now() WHERE id = %s", (opportunity_id,))
        log_event("materials_failed", f"{row['title']}: {exc}", level="warn", run_id=run_id, opportunity_id=opportunity_id)
        return {**summary, "outcome": "materials_failed", "error": str(exc)}

    verified = verification.all_claims_supported and not verification.unsupported_claims
    version = int(db.fetch_value("SELECT COALESCE(MAX(version), 0) AS v FROM application_materials WHERE opportunity_id = %s", (opportunity_id,)) or 0) + 1
    output_dir = write_files(row, materials, verification, version, cfg)

    with db.connection() as conn:
        db.execute(
            """
            INSERT INTO application_materials (opportunity_id, version, cv_markdown, cover_letter, verification, verified, model, output_dir)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (opportunity_id, version, materials.cv_markdown, materials.cover_letter_markdown,
             db.jsonb({**verification.model_dump(), "tailoring_notes": materials.tailoring_notes}), verified, ai.model, output_dir),
            conn,
        )
        db.execute(
            "UPDATE opportunities SET materials_status = %s, updated_at = now() WHERE id = %s",
            ("generated" if verified else "needs_review", opportunity_id),
            conn,
        )
        log_event(
            "materials_generated" if verified else "materials_need_review",
            f"{row['title']}: materials v{version} " + ("verified against CV" if verified else "have unsupported claims - review before use"),
            level="info" if verified else "warn", run_id=run_id, opportunity_id=opportunity_id, conn=conn,
        )
    return {**summary, "outcome": "materials_generated" if verified else "materials_need_review", "version": version}


def write_files(row: dict[str, Any], materials: TailoredMaterials, verification: MaterialsVerification, version: int, cfg: Config) -> str:
    folder = f"{date.today().isoformat()}_{row['id']}_{_slug(row.get('company_name'))}_{_slug(row['title'])}"
    relative = Path(str(cfg.get("materials.output_dir", "output/applications"))) / folder / f"v{version}"
    target = OUTPUT_ROOT / relative
    target.mkdir(parents=True, exist_ok=True)

    (target / "cv.md").write_text(materials.cv_markdown, encoding="utf-8")
    (target / "cover_letter.md").write_text(materials.cover_letter_markdown, encoding="utf-8")
    markdown_to_docx(materials.cv_markdown, target / "cv.docx")
    markdown_to_docx(materials.cover_letter_markdown, target / "cover_letter.docx")
    (target / "verification.json").write_text(
        json.dumps({**verification.model_dump(), "tailoring_notes": materials.tailoring_notes}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    (target / "job.md").write_text(
        "\n".join([
            f"# {row['title']}",
            f"**Company:** {row.get('company_name') or 'Not stated'}",
            f"**Posting:** {row['source_url']}",
            f"**Apply:** {row.get('apply_email') or row.get('apply_url') or 'See posting'} ({row.get('apply_method')})",
            f"**Match score:** {row.get('match_score')} - {row.get('recommendation')}",
            f"**Nepal eligibility:** {row.get('nepal_eligibility')} - {row.get('nepal_evidence')}",
            "",
            "## Why",
            *[f"- {reason}" for reason in (row.get("recommendation_reasons") or [])],
            "",
            "## Posting",
            row.get("description") or "",
        ]),
        encoding="utf-8",
    )
    return relative.as_posix()


_BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _add_runs(paragraph, text: str) -> None:
    for part in _BOLD.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part.replace("*", ""))


def markdown_to_docx(markdown: str, path: Path | Any) -> None:
    """Write a simple, ATS-friendly .docx. `path` may be a file path or a binary stream."""
    document = Document()
    for line in markdown.splitlines():
        text = line.rstrip()
        stripped = text.lstrip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ", "• ")):
            _add_runs(document.add_paragraph(style="List Bullet"), stripped[2:])
        else:
            _add_runs(document.add_paragraph(), stripped)
    document.save(str(path) if isinstance(path, Path) else path)
